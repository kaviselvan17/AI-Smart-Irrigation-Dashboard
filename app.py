
import os
from flask import request, jsonify
from openai import OpenAI

from flask import Flask, render_template, request, jsonify, send_file
from datetime import datetime, timedelta
from threading import Lock
from io import BytesIO
import os
import logging

from ai_model.predict import predict_irrigation
from voice_ai.assistant import get_voice_response

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment

from database import (
    initialize_database,
    save_sensor_reading,
    get_sensor_history,
    save_pump_log
)


# ==========================================
# FLASK APPLICATION
# ==========================================

app = Flask(__name__)
app.config["JSON_SORT_KEYS"] = False

logging.basicConfig(level=logging.INFO)

DEVICE_TOKEN = os.environ.get("DEVICE_TOKEN", "").strip()

data_lock = Lock()


# ==========================================
# SENSOR DATA
# ==========================================

sensor_data = {
    "soil1_raw": 0,
    "soil2_raw": 0,
    "soil1_moisture": 0.0,
    "soil2_moisture": 0.0,
    "rain_raw": 0,
    "rain_detected": False,
    "flow_pulses": 0,
    "water_liters": 0.0,
    "pump_status": False,
    "flow_fault": False,
    "runtime_fault": False,
    "automatic_mode": True,
    "last_update": None,
    "last_seen": None
}

control_state = {
    "mode": "AUTO",
    "pump": 0
}


# ==========================================
# AUTHORIZATION
# ==========================================

def authorized():
    """Check the ESP32 bearer token if one is configured."""

    if not DEVICE_TOKEN:
        return True

    auth_header = request.headers.get("Authorization", "")
    expected = f"Bearer {DEVICE_TOKEN}"

    return auth_header == expected


# ==========================================
# HELPER FUNCTIONS
# ==========================================

def get_device_online(last_seen):
    """Consider the ESP32 online if data arrived recently."""

    if not last_seen:
        return False

    try:
        last_time = datetime.fromisoformat(last_seen)
        return datetime.now() - last_time < timedelta(seconds=10)
    except (ValueError, TypeError):
        return False


def get_sensor_snapshot():
    """Return a thread-safe copy of the current sensor data."""

    with data_lock:
        data = sensor_data.copy()
        current_control = control_state.copy()

    data["device_online"] = get_device_online(data["last_seen"])
    data["mode"] = current_control["mode"]

    return data


# ==========================================
# HOME PAGE
# ==========================================

@app.route("/")
def index():
    return render_template("index.html")


# ==========================================
# HEALTH CHECK
# ==========================================

@app.route("/api/health", methods=["GET"])
def health():
    return jsonify({
        "status": "online",
        "application": "Smart Irrigation Dashboard",
        "time": datetime.now().isoformat()
    })


# ==========================================
# RECEIVE SENSOR DATA FROM ESP32
# ==========================================

@app.route("/api/sensors", methods=["POST"])
def receive_sensor_data():

    if not authorized():
        return jsonify({
            "error": "Unauthorized device"
        }), 401

    payload = request.get_json(silent=True)

    if not isinstance(payload, dict):
        return jsonify({
            "error": "Expected a JSON object"
        }), 400

    required_fields = [
        "soil1_raw",
        "soil2_raw",
        "soil1_moisture",
        "soil2_moisture",
        "rain_raw",
        "rain_detected",
        "flow_pulses",
        "water_liters",
        "pump_status",
        "flow_fault",
        "runtime_fault"
    ]

    missing_fields = [
        field for field in required_fields
        if field not in payload
    ]

    if missing_fields:
        return jsonify({
            "error": "Missing required fields",
            "fields": missing_fields
        }), 400

    try:
        soil1_raw = int(payload["soil1_raw"])
        soil2_raw = int(payload["soil2_raw"])
        soil1 = float(payload["soil1_moisture"])
        soil2 = float(payload["soil2_moisture"])
        rain_raw = int(payload["rain_raw"])
        flow_pulses = int(payload["flow_pulses"])
        water_liters = float(payload["water_liters"])

        if not 0 <= soil1 <= 100:
            raise ValueError("Soil sensor 1 must be between 0 and 100")

        if not 0 <= soil2 <= 100:
            raise ValueError("Soil sensor 2 must be between 0 and 100")

        if flow_pulses < 0 or water_liters < 0:
            raise ValueError("Flow values cannot be negative")

        clean_data = {
            "soil1_raw": soil1_raw,
            "soil2_raw": soil2_raw,
            "soil1_moisture": soil1,
            "soil2_moisture": soil2,
            "rain_raw": rain_raw,
            "rain_detected": bool(payload["rain_detected"]),
            "flow_pulses": flow_pulses,
            "water_liters": water_liters,
            "pump_status": bool(payload["pump_status"]),
            "flow_fault": bool(payload["flow_fault"]),
            "runtime_fault": bool(payload["runtime_fault"])
        }

    except (ValueError, TypeError, OverflowError) as error:
        return jsonify({
            "error": "Invalid sensor data",
            "details": str(error)
        }), 400

    now = datetime.now().isoformat()

    with data_lock:
        previous_pump_status = sensor_data["pump_status"]

        sensor_data.update(clean_data)
        sensor_data["last_update"] = now
        sensor_data["last_seen"] = now

        # Keep the dashboard's displayed operating mode updated.
        sensor_data["automatic_mode"] = (
            control_state["mode"] == "AUTO"
        )

        pump_changed = (
            previous_pump_status != clean_data["pump_status"]
        )

    # Save sensor readings to the database.
    try:
        save_sensor_reading(clean_data)
    except Exception:
        app.logger.exception("Could not save sensor reading to database")

    # Record pump state changes.
    # Adjust the arguments here if your database.py defines a
    # different save_pump_log() function signature.
    if pump_changed:
        try:
            save_pump_log(clean_data)
        except Exception:
            app.logger.exception("Could not save pump log")

    return jsonify({
        "status": "success",
        "message": "Sensor data received",
        "timestamp": now
    }), 200


# ==========================================
# GET CURRENT SENSOR DATA
# ==========================================

@app.route("/api/sensors", methods=["GET"])
def get_sensors():

    data = get_sensor_snapshot()

    return jsonify(data)


# ==========================================
# AI IRRIGATION PREDICTION
# ==========================================

@app.route("/api/ai-prediction", methods=["GET"])
def ai_prediction():

    data = get_sensor_snapshot()

    if not data["last_seen"]:
        return jsonify({
            "error": "Waiting for ESP32 sensor data"
        }), 503

    if not data["device_online"]:
        return jsonify({
            "error": "ESP32 is offline. Waiting for fresh sensor data."
        }), 503

    # Hardware safety rules take priority over AI predictions.
    if data["rain_detected"]:
        return jsonify({
            "irrigation_needed": False,
            "recommendation": "Do not irrigate while rain is detected.",
            "reason": "Rain detected. Safety rule overrides AI.",
            "confidence": None
        })

    if data["flow_fault"] or data["runtime_fault"]:
        return jsonify({
            "irrigation_needed": False,
            "recommendation": "Inspect the pump and resolve the safety fault.",
            "reason": "A pump safety fault is active.",
            "confidence": None
        })

    try:
        # IMPORTANT:
        # These argument names match your predict.py function:
        # predict_irrigation(soil1, soil2, rain, water_flow=0)
        result = predict_irrigation(
            soil1=data["soil1_moisture"],
            soil2=data["soil2_moisture"],
            rain=int(data["rain_detected"]),
            water_flow=data["water_liters"]
        )

        # Support prediction functions that return a dictionary.
        if isinstance(result, dict):
            result.setdefault(
                "recommendation",
                "Irrigation may be needed."
                if result.get("irrigation_needed")
                else "Irrigation is not recommended."
            )
            result.setdefault(
                "reason",
                "Prediction based on the current sensor readings."
            )
            result.setdefault("confidence", None)

            return jsonify(result)

        # Support prediction functions that return a boolean.
        if isinstance(result, bool):
            return jsonify({
                "irrigation_needed": result,
                "recommendation": (
                    "Irrigation may be needed."
                    if result
                    else "Irrigation is not recommended."
                ),
                "reason": "Based on the AI model prediction.",
                "confidence": None
            })

        app.logger.error(
            "Unexpected prediction result type: %s",
            type(result).__name__
        )

        return jsonify({
            "error": "Unexpected response from the AI prediction function"
        }), 500

    except FileNotFoundError:
        app.logger.exception("AI model file not found")
        return jsonify({
            "error": "AI model not found. Train the model first."
        }), 500

    except Exception:
        app.logger.exception("AI prediction failed")
        return jsonify({
            "error": "AI prediction failed. Check predict.py and server logs."
        }), 500


# ==========================================
# GET PUMP CONTROL COMMAND
# ==========================================

@app.route("/api/control", methods=["GET"])
def get_control():

    if not authorized():
        return jsonify({
            "error": "Unauthorized device"
        }), 401

    with data_lock:
        current_control = control_state.copy()

    return jsonify(current_control)


# ==========================================
# SET PUMP CONTROL MODE
# ==========================================

@app.route("/api/control", methods=["POST"])
def set_control():

    if not authorized():
        return jsonify({
            "error": "Unauthorized request"
        }), 401

    payload = request.get_json(silent=True)

    if not isinstance(payload, dict):
        return jsonify({
            "error": "Expected a JSON object"
        }), 400

    mode = str(payload.get("mode", "")).upper()

    if mode not in ("AUTO", "MANUAL"):
        return jsonify({
            "error": "Mode must be AUTO or MANUAL"
        }), 400

    # AUTO mode lets the ESP32 make the irrigation decision.
    # A zero command means no manual ON request.
    requested_pump = payload.get("pump", 0)

    if requested_pump not in (0, 1, False, True):
        return jsonify({
            "error": "Pump must be 0 or 1"
        }), 400

    pump = int(requested_pump)

    with data_lock:
        if mode == "AUTO":
            pump = 0

        control_state["mode"] = mode
        control_state["pump"] = pump

        sensor_data["automatic_mode"] = (mode == "AUTO")

        result_control = control_state.copy()

    # Log mode/command changes.
    # This call uses the same dictionary-based interface as above.
    try:
        save_pump_log({
            "pump_status": bool(pump),
            "mode": mode,
            "source": "dashboard",
            "timestamp": datetime.now().isoformat()
        })
    except Exception:
        app.logger.exception("Could not save dashboard pump command")

    return jsonify({
        "status": "success",
        "success": True,
        "message": "Control command updated",
        "control": result_control,
        "command": result_control
    }), 200


# ==========================================
# DOWNLOAD SENSOR HISTORY AS EXCEL
# ==========================================

@app.route("/download-excel", methods=["GET"])
def download_excel():

    try:
        history = get_sensor_history()

        if history is None:
            history = []

        workbook = Workbook()
        sheet = workbook.active
        sheet.title = "Sensor History"

        headers = [
            "Soil 1 Raw",
            "Soil 2 Raw",
            "Soil 1 Moisture (%)",
            "Soil 2 Moisture (%)",
            "Rain Raw",
            "Rain Detected",
            "Flow Pulses",
            "Water Liters",
            "Pump Status"
        ]

        sheet.append(headers)

        # Style the header row.
        for cell in sheet[1]:
            cell.font = Font(bold=True, color="FFFFFF")
            cell.fill = PatternFill(
                fill_type="solid",
                fgColor="167D5A"
            )
            cell.alignment = Alignment(horizontal="center")

        # Accept either dictionaries or tuples from the database.
        if isinstance(history, list):
            for record in history:
                if isinstance(record, dict):
                    sheet.append([
                        record.get("soil1_raw", ""),
                        record.get("soil2_raw", ""),
                        record.get("soil1_moisture", ""),
                        record.get("soil2_moisture", ""),
                        record.get("rain_raw", ""),
                        record.get("rain_detected", ""),
                        record.get("flow_pulses", ""),
                        record.get("water_liters", ""),
                        record.get("pump_status", "")
                    ])
                elif isinstance(record, (tuple, list)):
                    sheet.append(list(record))

        # Fit column widths.
        for column_cells in sheet.columns:
            max_length = max(
                len(str(cell.value or ""))
                for cell in column_cells
            )
            sheet.column_dimensions[
                column_cells[0].column_letter
            ].width = min(max_length + 2, 28)

        output = BytesIO()
        workbook.save(output)
        output.seek(0)

        filename = (
            "smart_irrigation_"
            + datetime.now().strftime("%Y%m%d_%H%M%S")
            + ".xlsx"
        )

        return send_file(
            output,
            as_attachment=True,
            download_name=filename,
            mimetype=(
                "application/vnd.openxmlformats-officedocument."
                "spreadsheetml.sheet"
            )
        )

    except Exception:
        app.logger.exception("Excel export failed")
        return jsonify({
            "error": "Could not export sensor history. Check database.py."
        }), 500

# ==========================================
# VOICE AI API
# ==========================================

@app.route("/api/voice", methods=["POST"])
def voice_api():
    payload = request.get_json(silent=True) or {}
    question = payload.get("question", "")

    if not isinstance(question, str) or not question.strip():
        return jsonify({"error": "Please provide a question."}), 400

    # Read the real sensor values from Flask.
    data = get_sensor_snapshot()

    # Pass the existing AI prediction to the voice assistant.
    prediction = None

    if data.get("last_seen") and data.get("device_online"):
        if data.get("rain_detected"):
            prediction = {
                "irrigation_needed": False,
                "recommendation": "Do not irrigate while rain is detected."
            }
        elif data.get("flow_fault") or data.get("runtime_fault"):
            prediction = {
                "irrigation_needed": False,
                "recommendation": "Check the pump safety fault."
            }
        else:
            try:
                prediction = predict_irrigation(
                    soil1=data["soil1_moisture"],
                    soil2=data["soil2_moisture"],
                    rain=int(data["rain_detected"]),
                    water_flow=data["water_liters"]
                )
            except Exception:
                app.logger.exception("Voice AI prediction failed")

    # Give the assistant the same field names it expects.
    voice_sensor_data = {
        **data,
        "soil_moisture": data.get("soil1_moisture"),
        "soil_moisture_2": data.get("soil2_moisture"),
        "rain": data.get("rain_detected"),
        "flow": data.get("water_liters"),
        "pump": "ON" if data.get("pump_status") else "OFF"
    }

    try:
        answer = get_voice_response(
            question,
            sensor_data=voice_sensor_data,
            ai_prediction=prediction
        )

        return jsonify({"answer": answer}), 200

    except Exception:
        app.logger.exception("Voice assistant failed")
        return jsonify({
            "error": "Voice assistant failed. Check the Flask terminal."
        }), 500

@app.route("/api/transcribe", methods=["POST"])
def transcribe_audio():
    if not os.environ.get("OPENAI_API_KEY"):
        return jsonify({"error": "Speech-to-text API key is not configured."}), 500

    if "audio" not in request.files:
        return jsonify({"error": "No audio file received."}), 400

    audio = request.files["audio"]

    if not audio.filename:
        return jsonify({"error": "Empty audio recording."}), 400

    try:
        client = OpenAI(api_key=os.environ["OPENAI_API_KEY"])

        result = client.audio.transcriptions.create(
            model="gpt-4o-mini-transcribe",
            file=(audio.filename, audio.stream, audio.mimetype)
        )

        return jsonify({"text": result.text})

    except Exception:
        app.logger.exception("Audio transcription failed")
        return jsonify({
            "error": "Transcription failed. Check server logs and API configuration."
        }), 500


# ==========================================
# START APPLICATION
# ==========================================

if __name__ == "__main__":

    try:
        initialize_database()
        app.logger.info("Database initialized")
    except Exception:
        app.logger.exception("Database initialization failed")
        raise

    app.run(
        host="0.0.0.0",
        port=5000,
        debug=False
    )
